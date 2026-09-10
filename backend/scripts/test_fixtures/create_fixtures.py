"""Create test fixture documents (TXT, DOCX, PDF) for upload testing.

Run: python scripts/test_fixtures/create_fixtures.py
"""

import os
import io
from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent

SPEC_TEXT = """PROCUREMENT SPECIFICATION FOR AIRCRAFT WOVEN CARPET

1. SCOPE
This specification covers the supply and installation of aircraft woven carpet suitable for aircraft interior applications.

2. MATERIAL REQUIREMENTS
The carpet shall be manufactured from textile floor covering construction. The pile shall be woven and the carpet shall meet the requirements for durability, dimensional stability, and resistance to abrasion.

3. PERFORMANCE REQUIREMENTS
The carpet shall comply with applicable fire/safety performance requirements for aircraft interior applications. It shall be tested for flame resistance, smoke emission, and toxicity in accordance with the relevant Indian Standard specification for aircraft woven carpet.

4. DIMENSIONS
The carpet shall be supplied in rolls with the following nominal dimensions:
Width: 1520 mm nominal
Thickness: 6.0 mm nominal
Mass per unit area: 1800 g/m2

5. TESTING
The carpet shall be tested in accordance with the applicable testing requirements, including tests for dimensional stability, abrasion resistance, and fire performance. All testing shall be carried out by an accredited laboratory.

6. MARKING AND PACKING
Each roll shall be marked with the standard number, batch number, and date of manufacture.
"""


def create_txt() -> Path:
    path = FIXTURES_DIR / "procurement_spec.txt"
    path.write_text(SPEC_TEXT, encoding="utf-8")
    return path


def create_docx() -> Path:
    """DOCX with headings, paragraphs, and a specification table."""
    import docx
    from docx.shared import Pt

    path = FIXTURES_DIR / "procurement_spec.docx"
    doc = docx.Document()

    doc.add_heading("PROCUREMENT SPECIFICATION FOR AIRCRAFT WOVEN CARPET", level=1)

    doc.add_heading("1. SCOPE", level=2)
    doc.add_paragraph(
        "This specification covers the supply and installation of aircraft woven carpet "
        "suitable for aircraft interior applications."
    )

    doc.add_heading("2. MATERIAL REQUIREMENTS", level=2)
    doc.add_paragraph(
        "The carpet shall be manufactured from textile floor covering construction. "
        "The pile shall be woven and the carpet shall meet the requirements for durability, "
        "dimensional stability, and resistance to abrasion."
    )

    doc.add_heading("3. PERFORMANCE REQUIREMENTS", level=2)
    doc.add_paragraph(
        "The carpet shall comply with applicable fire/safety performance requirements for "
        "aircraft interior applications. It shall be tested for flame resistance, smoke "
        "emission, and toxicity in accordance with the relevant Indian Standard specification "
        "for aircraft woven carpet."
    )

    doc.add_heading("4. DIMENSIONS", level=2)
    doc.add_paragraph("The carpet shall be supplied in rolls with the following nominal dimensions:")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Light Shading Accent 1"
    data = [
        ("Property", "Requirement"),
        ("Width", "1520 mm nominal"),
        ("Thickness", "6.0 mm nominal"),
        ("Mass per unit area", "1800 g/m2"),
    ]
    for i, (col1, col2) in enumerate(data):
        row = table.rows[i]
        row.cells[0].text = col1
        row.cells[1].text = col2

    doc.add_heading("5. TESTING", level=2)
    doc.add_paragraph(
        "The carpet shall be tested in accordance with the applicable testing requirements, "
        "including tests for dimensional stability, abrasion resistance, and fire performance. "
        "All testing shall be carried out by an accredited laboratory."
    )

    doc.add_heading("6. MARKING AND PACKING", level=2)
    doc.add_paragraph(
        "Each roll shall be marked with the standard number, batch number, and date of manufacture."
    )

    doc.save(path)
    return path


def create_pdf() -> Path:
    """Text-based PDF using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors

    path = FIXTURES_DIR / "procurement_spec.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleX", parent=styles["Title"], fontSize=15, leading=19, spaceAfter=12
    )
    heading_style = ParagraphStyle(
        "HeadingX", parent=styles["Heading2"], fontSize=13, leading=16, spaceAfter=6
    )
    body_style = styles["BodyText"]

    story = []
    story.append(Paragraph("PROCUREMENT SPECIFICATION FOR AIRCRAFT WOVEN CARPET", title_style))

    story.append(Paragraph("1. SCOPE", heading_style))
    story.append(Paragraph(
        "This specification covers the supply and installation of aircraft woven carpet "
        "suitable for aircraft interior applications.", body_style))

    story.append(Paragraph("2. MATERIAL REQUIREMENTS", heading_style))
    story.append(Paragraph(
        "The carpet shall be manufactured from textile floor covering construction. The pile "
        "shall be woven and the carpet shall meet the requirements for durability, dimensional "
        "stability, and resistance to abrasion.", body_style))

    story.append(Paragraph("3. PERFORMANCE REQUIREMENTS", heading_style))
    story.append(Paragraph(
        "The carpet shall comply with applicable fire/safety performance requirements for "
        "aircraft interior applications. It shall be tested for flame resistance, smoke "
        "emission, and toxicity in accordance with the relevant Indian Standard specification "
        "for aircraft woven carpet.", body_style))

    story.append(Paragraph("4. DIMENSIONS", heading_style))
    story.append(Paragraph("The carpet shall be supplied in rolls with the following nominal dimensions:", body_style))

    table_data = [
        ["Property", "Requirement"],
        ["Width", "1520 mm nominal"],
        ["Thickness", "6.0 mm nominal"],
        ["Mass per unit area", "1800 g/m2"],
    ]
    tbl = Table(table_data, colWidths=[60 * mm, 90 * mm])
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ]))
    story.append(Spacer(1, 6))
    story.append(tbl)
    story.append(Spacer(1, 6))

    story.append(Paragraph("5. TESTING", heading_style))
    story.append(Paragraph(
        "The carpet shall be tested in accordance with the applicable testing requirements, "
        "including tests for dimensional stability, abrasion resistance, and fire performance. "
        "All testing shall be carried out by an accredited laboratory.", body_style))

    story.append(Paragraph("6. MARKING AND PACKING", heading_style))
    story.append(Paragraph(
        "Each roll shall be marked with the standard number, batch number, and date of manufacture.",
        body_style))

    doc.build(story)
    return path


if __name__ == "__main__":
    created = []
    created.append(create_txt())
    created.append(create_docx())
    created.append(create_pdf())
    for p in created:
        print(f"Created: {p} ({p.stat().st_size} bytes)")
    print("All fixtures created.")